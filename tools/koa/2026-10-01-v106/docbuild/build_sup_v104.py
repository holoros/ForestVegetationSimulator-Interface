# -*- coding: utf-8 -*-
"""Build koa_supplemental_v104.docx from s103 by anchored edits."""
import sys, copy, datetime, json
sys.path.insert(0, '/tmp/k6/work')
import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from lxml import etree
import docedit as E

SRC = '/tmp/k6/src/koa_supplemental_v103.docx'
OUT = sys.argv[1] if len(sys.argv) > 1 else '/tmp/k6/out/koa_supplemental_v104.docx'
d = docx.Document(SRC)
P = d.paragraphs
log = []

def plain(s):
    return E.expand(s).replace('^', '').replace('~', '').replace('**', '')

def sub(anchor, old, new, count=1):
    p = E.find_par(P, anchor, startswith=True)
    E.replace_in(p, plain(old), new, count)
    log.append(('SUB', anchor[:40], plain(old)[:50]))

def table_after(caption_start):
    p = E.find_par(P, caption_start, startswith=True)
    el = p._p.getnext()
    while el is not None and etree.QName(el).localname != 'tbl':
        el = el.getnext()
    return Table(el, p._parent)

def cell_sub(t, r, c, old, new):
    E.replace_in(t.rows[r].cells[c].paragraphs[0], plain(old), new)

def flip(s):
    s = s.strip()
    if s.startswith('+'): return '−' + s[1:]
    if s.startswith('−'): return '+' + s[1:]
    if s.startswith('-'): return '+' + s[1:]
    raise ValueError('unsigned value ' + s)

# ---------------- front matter ----------------
sub('This file supports the manuscript named above.', 'Tables S1 to S32 and Figures S1 to S8', 'Tables S1 to S33 and Figures S1 to S8')

# ---------------- S1.1 ----------------
sub('Tree data were compiled from eight monitoring programs',
    'Measurement intervals ranged from 1 to 16 years (median 4 years).',
    'Consecutive remeasurement intervals range from 1 to 23 years, with a median of 1 year on the consecutive diameter frame, where the annual PSP and KMR PSP visits dominate (median 5 years in DOFAW and 9 years in FIA).')
sub('Thinning removals are a second property of the PSP network',
    'Of the 158 installations, 151 have at least two measurements.',
    'Of the 158 installations, 70 carry measurements in two or more years (67 on live records), and every plot of the Kahikinui, KMR CAR, Kualoa, Mauka and Kap sources was measured once.')

# ---------------- S1.3 ----------------
sub('The repair rules were written into the job record',
    '22 had deposited values 1.9 to 2.1 times the rebuilt value, 48 exceeded it by more than 10% and 164',
    '22 had deposited values 1.9 to 2.1 times the rebuilt value, 32 more exceeded it by more than 10% (54 in total) and 164')
sub('The repair moved at least one covariate by more than 5%',
    'since they rebuild stand values from the live list.',
    'since they rebuild stand values from the live list. Table S33 reconciles every increment frame size quoted in the main text and this file, from the 13,122 consecutive tree intervals of the tree table to the frames on which the deployed equations were fitted and calibrated, and the 4,670 and 3,812 records above are the diameter and height keys common to the original and recomputed consecutive frames.')

# ---------------- S2.1 ----------------
sub('Aboveground biomass for each FIA subplot visit was computed',
    'are illustrative levels within the range of plot BYI in the growth data (0 to 813).',
    'are illustrative levels within the range of plot BYI in the growth data (0 to 813). One growth-model plot, Kulani 42, carries BYI 0, and its 30 live records enter only Eq. 2, where BYI acts linearly on the asymptote, since no increment record has BYI below 25 and no crown record below 130, so ln(BYI) is defined wherever it is used.')

# ---------------- S3.1, S3.2, S3.3 ----------------
sub('The deployed crown base vector, carried by the projection engine',
    'The preregistered rule required both a lower RMSE and every sign kept,',
    'The selection rule fixed before refitting required both a lower RMSE and every sign kept,')
sub('Preregistered selection. Before any increment was fitted',
    'Preregistered selection.', '**Selection rule.**')
sub('Back-transformation and calibration. The deployed diameter factor',
    'exp(0.5 × (0.586222 + 0.533222)) = 1.369',
    'exp(0.5 × (0.586222^2^ + 0.533222^2^)) = 1.369')
sub('Solved within each data source on the repaired calibration frame',
    'tested on twenty pre-registered installation splits',
    'tested on twenty installation splits fixed in advance')
sub('Re-solved constants. Holding the published diameter vector fixed',
    '4.7% and 1.6% below the calibration-frame values',
    '4.5% and 1.6% below the calibration-frame values')

# ---------------- S6.2 sign convention ----------------
sub('Because the level factor was solved on intervals from the installations',
    'The pooled natural survival error was −0.057 (RMSE 0.182) held out, against −0.058 (0.168) in sample and −0.268 (0.320) without the factor (observed minus projected; Table S28), and on the 23-unit validation the held-out factors moved the natural survival bias from −0.198 to −0.200.',
    'The pooled natural survival error was +0.057 (RMSE 0.182) held out, against +0.058 (0.168) in sample and +0.268 (0.320) without the factor (projected minus observed; Table S28), and on the 23-unit validation the held-out factors moved the natural survival bias from +0.198 to +0.200.')
sub('Stand-level calibration was tested under the earlier system',
    'On the four natural plots it removed most of the QMD bias (+3.53 to −1.06 cm, change −4.59 cm, 95% plot-cluster interval −6.39 to −2.50)',
    'On the four natural plots it removed most of the QMD bias (projected minus observed −3.53 to +1.06 cm, change +4.59 cm, 95% plot-cluster interval +2.50 to +6.39)')
sub('Stand-level calibration was tested under the earlier system',
    'it raised the QMD error by 2.21 cm (0.19 to 4.04) at one to five years',
    'it deepened the QMD underprojection by 2.21 cm (0.19 to 4.04) at one to five years')
sub('Data flags. Kulani plot 12 has a mapped BYI of 25.1',
    'which places their trajectory points near +0.99 in Fig. 5d,',
    'which places their survival errors at +0.99, and Fig. 5 marks them and excludes them from the horizon means,')

# ---------------- S7 ----------------
sub('Table S30 gives the full projection grid at ages 20, 40, 60 and 100 yr, and Figs. S6',
    'The natural stands therefore never reach half of the planted peak density or of the SDI observed on three DOFAW natural plots (623 to 924), an envelope the BAL correction did not change.',
    'The natural stands therefore reach about half of the planted peak density and stay well below the rebuilt plot-year SDI maxima of the three long natural DOFAW plots (762 to 880, or 623 to 924 when SDI is computed from QMD at the start of each interval), an envelope the BAL correction did not change.')
sub('Bakuzis matrix. The modified Bakuzis matrix',
    '−1.20, −1.29 and −1.32 for uneven-aged stands, so seven of nine combinations fall in the −1.2 to −2.2 band,',
    '−1.18, −1.29 and −1.32 for uneven-aged stands, so seven of nine combinations fall in the −1.2 to −2.2 band and the low natural sites (−1.02 even-aged, −1.18 uneven-aged) are shallower than it,')

sub('Each annual step (1) computes BAL by Eq. 7',
    'Ingrowth enters at 2.5 cm (Eq. 6), truncated',
    'Ingrowth enters at 2.5 cm (Eq. 6), whose coefficients refitted on the recomputed covariates lie within 0.06 standard errors of the original fit, truncated')

# ---------------- references ----------------
sub('[dataset] Weiskittel, A.R., Sprecher, I.', 'Version 1.9.3.', 'Version 1.9.5.')

# ---------------- captions ----------------
sub('Table S8. Fit statistics of the height equation', 'Bias is observed minus predicted.', 'Bias is predicted minus observed.')
sub('Table S16. Performance of the five koa component equations',
    'where in-sample origin bias is predicted minus observed.',
    'and every bias is predicted minus observed.')
sub('Table S18. ', 'than the 9,059-record frame of Eq. 2', 'than the 8,914-record frame of Eq. 2')
sub('Table S24. Monte Carlo parameter distributions',
    'Multiplier and Stage 1 ranges are 2.5th and 97.5th percentiles across rows,',
    'The draw spread column gives the standard deviation across the drawn rows (for the height parameters, against the model standard errors of Table 3), the standard deviation of log k for the multipliers and the bootstrap standard deviation for Stage 1, and multiplier and Stage 1 ranges are 2.5th and 97.5th percentiles across rows,')
sub('Table S28. ', 'Errors are observed minus projected cohort survival,', 'Errors are projected minus observed cohort survival,')
sub('Table S29. ', 'observed minus projected cohort survival,', 'projected minus observed cohort survival,')
sub('Table S29. ', 'Change is local minus population', 'Change is local minus population error')
sub('Figure S1. Diagnostics for the koa height', 'Height residuals carry a population-average bias of +0.39 m observed minus predicted,', 'Height predictions carry a population-average bias of −0.39 m predicted minus observed,')
sub('Figure S7. Even-aged planted koa projections',
    'Planted QMD approaches the 69.7 cm planted diameter ceiling late in the century,',
    'The largest planted tree reaches the 69.7 cm planted diameter ceiling by about age 40, whereas planted QMD tops out at 59.8 cm,')


# ---------------- round 2 (red team 5 and stress test 2) ----------------
sub('Aboveground biomass for each FIA subplot visit was computed', 'and no crown record below 130,', 'and no crown record below 129,')
sub('Table S30 gives the full projection grid at ages 20, 40, 60 and 100 yr, and Figs. S6',
    'stay well below the rebuilt plot-year SDI maxima of the three long natural DOFAW plots (762 to 880, or 623 to 924 when SDI is computed from QMD at the start of each interval), an envelope the BAL correction did not change.',
    'stay well below the rebuilt plot-year SDI maxima of 762 to 880 on three of the four long natural DOFAW plots, Kulani 23, Laupahoehoe 41 and Waikamoi 25 (623 to 924 when SDI is computed from QMD at the start of each interval), although Waiakea 24 reached only 399, an envelope the BAL correction did not change.')
sub('This file supports the manuscript named above.',
    'Results labeled as from the earlier system come from the deposited covariates and the engine of the previous revision and were not repeated on the repaired frames, and the caption of every such table says so.',
    'Results marked as on the original covariates use the stand covariates as first assembled, before the recomputation of Section S1.3, and the simulator before its weighted competition update, and were not repeated on the recomputed frames, and the caption of every such table says so.')
for p_ in P:
    if p_.text.startswith('Table S') and 'Carried from the earlier system on the deposited covariates and not rerun on the repaired frames.' in p_.text:
        E.replace_in(p_, 'Carried from the earlier system on the deposited covariates and not rerun on the repaired frames.', 'On the original covariates, not repeated on the recomputed frames.')
        log.append(('CAPLABEL', p_.text[:10]))
sub('Selection rule. Before any increment was fitted', 'and with the smallest largest equivalence region as the fallback', 'and with the candidate whose widest equivalence region was narrowest as the fallback')
sub('Selection rule. Before any increment was fitted', 'b6 < 0; height', 'b6 < 0, and height')
sub('Table S27. Trajectory validation errors',
    'Planted rows carry a single value because the factor applies only to natural stands.',
    'Planted rows carry a single value because the factor applies only to natural stands. The planted 1 to 5 yr row includes the four PSP 105 to 108 visits of 2009 that the source data code as entirely dead, and without them, as in the horizon means of Fig. 5, the planted biases at 1 to 5 yr are +0.057 for survival, −1.59 cm for QMD and −2.7 {m2ha} for basal area.')
t = table_after('Table S2. Summary statistics of the koa data by source')
rows = {t.rows[i].cells[0].text: i for i in range(len(t.rows))}
assert t.rows[rows['Total']].cells[5].text == '19.2'
E.set_par(t.rows[rows['Total']].cells[5].paragraphs[0], '19.5')
t = table_after('Table S16. Performance of the five koa component equations')
for i in range(1, len(t.rows)):
    if t.rows[i].cells[2].text == 'AUC, leave-one-installation-out':
        E.set_par(t.rows[i].cells[2].paragraphs[0], 'AUC, leave-one-installation-out, pooled interval scale')

# ---------------- table cells ----------------
t = table_after('Table S2. Summary statistics of the koa data by source')
rows = {t.rows[i].cells[0].text: i for i in range(len(t.rows))}
E.set_par(t.rows[rows['Kahikinui']].cells[4].paragraphs[0], '5.3')

# S8 height and HCB bias rows (to predicted minus observed)
t = table_after('Table S8. Fit statistics of the height equation')
rows = {t.rows[i].cells[0].text: i for i in range(len(t.rows))}
for k in ['Conditional bias (m)', 'Population-average bias (m)']:
    for c in (1, 2):
        cell = t.rows[rows[k]].cells[c].paragraphs[0]
        E.set_par(cell, flip(cell.text))
k = 'Population-average bias, natural and planted (m)'
E.set_par(t.rows[rows[k]].cells[1].paragraphs[0], '−0.232 and −1.003')
E.set_par(t.rows[rows[k]].cells[2].paragraphs[0], '+0.008 and n/a')
E.set_par(t.rows[rows['Leave-one-installation-out bias (m), population-average']].cells[1].paragraphs[0], '−0.053')
k = 'Equivalence, mean bias (90% interval) against region'
E.set_par(t.rows[rows[k]].cells[1].paragraphs[0], '−0.395 (−0.681 to −0.101) against ±2.29, equivalent')

# S16 height and HCB bias rows
t = table_after('Table S16. Performance of the five koa component equations')
for i in range(1, len(t.rows)):
    comp, stat = t.rows[i].cells[0].text, t.rows[i].cells[2].text
    if comp in ('Height', 'HCB') and stat.startswith('bias'):
        for c in (3, 4):
            p = t.rows[i].cells[c].paragraphs[0]
            if p.text.strip() == 'n/a':
                continue
            parts = p.text.split(' (')
            new = flip(parts[0]) + ((' (' + flip(parts[1].rstrip(')')) + ')') if len(parts) > 1 else '')
            E.set_par(p, new)
            log.append(('S16FLIP', comp, stat, p.text))

# S17 note and koa row
t = table_after('Table S17. Performance of the A. koa equations')
for i in range(len(t.rows)):
    for c in t.rows[i].cells:
        for p in c.paragraphs:
            if 'whereas the deployed population-average values are 0.786 (height) and 0.317 (calibrated ΔDBH)' in p.text:
                E.replace_in(p, 'whereas the deployed population-average values are 0.786 (height) and 0.317 (calibrated ΔDBH)',
                             'whereas the deployed population-average values are 0.788 (height) and 0.381 (calibrated ΔDBH)')
                E.replace_in(p, 'and the survival AUC is apparent,', 'and the survival AUC is the apparent value of the published equation, which is not deployed,')
                log.append(('S17NOTE',))
    if t.rows[i].cells[0].text == 'Tectona grandis':
        E.set_par(t.rows[i].cells[0].paragraphs[0], '*Tectona grandis* L.f.')
    if t.rows[i].cells[0].text == 'Eucalyptus grandis':
        E.set_par(t.rows[i].cells[0].paragraphs[0], '*Eucalyptus grandis* W.Hill ex Maiden')

# S24 header
t = table_after('Table S24. Monte Carlo parameter distributions')
assert t.rows[0].cells[2].text == 'SE'
E.set_par(t.rows[0].cells[2].paragraphs[0], '**Draw spread**')

# S28 and S29: flip bias columns
t = table_after('Table S28. ')
hdr = [c.text for c in t.rows[0].cells]
bcols = [j for j, h in enumerate(hdr) if 'bias' in h]
for i in range(1, len(t.rows)):
    for j in bcols:
        p = t.rows[i].cells[j].paragraphs[0]
        E.set_par(p, flip(p.text))
log.append(('S28FLIP', bcols))
t = table_after('Table S29. ')
hdr = [c.text for c in t.rows[0].cells]
bcols = [j for j, h in enumerate(hdr) if 'bias' in h]
for i in range(1, len(t.rows)):
    for j in bcols:
        p = t.rows[i].cells[j].paragraphs[0]
        txt = p.text
        if ' / ' in txt:
            a, b = txt.split(' / ')
            E.set_par(p, flip(a) + ' / ' + flip(b))
        else:
            head = txt.split(' (')[0]
            fl = flip(head) if head[0] in '+−-' else '−' + head
            if ' (' in txt:
                lo, hi = txt.split(' (')[1].rstrip(')').split(' to ')
                neg = lambda v: ('+' + v[1:] if v.startswith('−') else ('−' + v if v not in ('0', '0.00') else v))
                fl += ' (' + neg(hi) + ' to ' + neg(lo) + ')'
            E.set_par(p, fl)
log.append(('S29FLIP', bcols))

# S31 rounding
t = table_after('Table S31. Behavior scenarios')
for i in range(len(t.rows)):
    if t.rows[i].cells[0].text in ('Density, natural', 'Thinning, natural') and t.rows[i].cells[1].text in ('500 stems', 'Unthinned'):
        assert t.rows[i].cells[2].text == '34.7'
        E.set_par(t.rows[i].cells[2].paragraphs[0], '34.8')


# ---------------- round 3 (open items, October 1, 2026) ----------------
sub('The archived survival table (AK_SURV.csv) was built', 'and 7 of 62 folds fail,', 'and 7 of its 62 leave-one-installation-out refits fail to converge,')
sub('Reproducibility forensics are a companion document', '(5 of 62 jackknife folds failed)', '(5 of 62 delete-one-installation jackknife folds failed on the thinning-censored baseline sample of 4,584 records, a different procedure from the 7 failed leave-one-installation-out refits of Section S3.4)')
sub('Because the level factor was solved on intervals from the installations', ' (projected minus observed; Table S28), and on the 23-unit validation the held-out factors moved the natural survival bias from +0.198 to +0.200.', ' (projected minus observed; Table S28).')
sub('Three properties that a site index should have', '(nlme 3.1.170)', '(nlme 3.1-170)')
sub('Table S8 gives fit statistics for the height equation of Table 3',
    'so the negative relative-diameter term lowers the predicted height of the very largest trees on big-tree plots.',
    'so the negative relative-diameter term lowers the predicted height of the very largest trees on big-tree plots. Eq. 2 was fitted with rDBH recomputed from the live-tree list (tree_join_v103.csv and height_refit_v103.R in the deposit) as DBH over the largest DBH among live stems with a positive expansion factor in the plot-year. The DBH.max column carried in AK_TREE.csv and the static height frame is not that predictor, since in 73 of 739 plot-years it exceeds the largest live stem (60 values left from before deduplication and 13 dead stems), and refitting Eq. 2 with it returns a~0~ = 25.47 and a~1~ = 0.604.')

# ---------------- new Table S33 ----------------
cap32 = E.find_par(P, 'Table S32. Net mean annual increment', startswith=True)
t32 = table_after('Table S32. Net mean annual increment')
t21 = table_after('Table S21. ')
newcap = copy.deepcopy(cap32._p)
newtbl = copy.deepcopy(t21._tbl)
t32._tbl.addnext(newcap)
newcap.addnext(newtbl)
capP = Paragraph(newcap, cap32._parent)
E.set_par(capP, "**Table S33.** Reconciliation of the koa increment frames, from every consecutive tree interval of the tree table to the frames on which the diameter (ΔDBH) and height (ΔHT) increment equations were fitted and calibrated. Counts are tree intervals on the recomputed stand covariates unless stated, and the main text, Table 2, Table S12, Table S16 and Fig. S2 use the calibration frame. Because a first-to-last interval overlaps the consecutive intervals of the same tree, 179 diameter and 176 height trees contribute their growth twice to the calibration frame.")
T = Table(newtbl, cap32._parent)
rowsdef = [
 ('**Frame**', '**ΔDBH**', '**ΔHT**', '**Definition**'),
 ('All consecutive intervals', '13,122', '13,122', 'Consecutive tree intervals of the tree table, any status'),
 ('Live to live', '5,420', '5,420', 'Live with DBH > 0 at both ends, before the increment screens'),
 ('Screened', '4,813', '3,868', 'Annual increment above 0 and below 10 {cmyr} or 10 {myr}'),
 ('Consecutive frame', '4,792', '3,857', 'Screened intervals with a plot record, which drops 21 and 11 FIA intervals'),
 ('Calibration frame', '4,971', '4,033', 'Consecutive frame plus the 179 and 176 first-to-last intervals that span no thinning removal, with a recorded origin'),
 ('Calibration frame without origin screen', '4,992', '4,044', 'Adds 21 and 11 FIA intervals with no origin or BYI, not used in any fit'),
 ('Deployed fit', '4,790', '4,033', 'ΔDBH vector fitted to the consecutive frame on the original covariates, ΔHT vector fitted to the calibration frame'),
 ('Common to original and recomputed frames', '4,670', '3,812', 'Consecutive-frame keys present in both, used for the covariate change of Section S1.3'),
]
trs = T._tbl.findall(qn('w:tr'))
body_tpl = trs[1]
last_tpl = trs[-1]
for tr in trs[1:]:
    T._tbl.remove(tr)
for k, vals in enumerate(rowsdef[1:]):
    tpl = last_tpl if k == len(rowsdef) - 2 else body_tpl
    T._tbl.append(copy.deepcopy(tpl))
T = Table(newtbl, cap32._parent)
for i, vals in enumerate(rowsdef):
    for j, v in enumerate(vals):
        E.set_par(T.rows[i].cells[j].paragraphs[0], v)
W = [2500, 1000, 1000, 4860]
g = T._tbl.find(qn('w:tblGrid'))
for gc, w in zip(g, W):
    gc.set(qn('w:w'), str(w))
for row in T.rows:
    for cell, w in zip(row.cells, W):
        tcW = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
        tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
log.append(('S33', len(T.rows)))

# ---------------- superscript sweep (not references) ----------------
inrefs, nfix = False, 0
for p in P:
    if p.text.strip() == 'References': inrefs = True
    if p.text.strip() == 'Tables': inrefs = False
    if not inrefs:
        try: nfix += E.fix_superscripts(p)
        except ValueError as e: log.append(('SKIP', str(e)[:50]))
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                try: nfix += E.fix_superscripts(p)
                except ValueError as e: log.append(('SKIP', str(e)[:50]))
log.append(('SUPFIX', nfix))

cp = d.core_properties
cp.modified = datetime.datetime(2026, 10, 1, 9, 0, 0)
cp.revision = (cp.revision or 0) + 1
d.save(OUT)
json.dump(log, open(OUT.replace('.docx', '_buildlog.json'), 'w'), indent=0, ensure_ascii=False)
print('ok', len(log), 'superscripts fixed', nfix)
