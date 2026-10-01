# -*- coding: utf-8 -*-
"""BLIND copies of v106 and s104, same rules as the v105/s103 BLIND files."""
import sys, re
sys.path.insert(0, '/tmp/k6/work')
import docx
import docedit as E

def blank(p):
    E.set_par(p, '')

def run(src, out, kind):
    d = docx.Document(src)
    body = d.element.body
    for tag in ('commentRangeStart', 'commentRangeEnd'):
        for el in list(body.iter('{%s}%s' % (E.W, tag))):
            el.getparent().remove(el)
    for el in list(body.iter('{%s}commentReference' % E.W)):
        r = el.getparent(); r.getparent().remove(r)
    for rel in list(d.part.rels.values()):
        if rel.reltype.endswith('/comments') or 'commentsExtended' in rel.reltype or 'commentsIds' in rel.reltype:
            d.part.drop_rel(rel.rId)
    P = d.paragraphs
    if kind == 'ms':
        E.set_par(E.find_par(P, 'Aaron Weiskittela', startswith=True), 'Authors and affiliations withheld for double-blind review.')
        for p in P:
            t = p.text.strip()
            if len(t) < 120 and t.startswith(('a Center for Research', 'b Forest Solutions', 'c Midgard', '*Corresponding author', 'Center for Research on Sustainable Forests, University of Maine')):
                blank(p)
        for a in ['We thank the Hawaii Division', 'Aaron Weiskittel: Conceptualization', 'This work was supported by the USDA']:
            E.set_par(E.find_par(P, a, startswith=True), 'Withheld for double-blind review.')
        p = E.find_par(P, 'The FVS-HI variant (HiGy.R,', startswith=True)
        E.replace_in(p, 'https://github.com/MidgardNaturalResources/ForestVegetationSimulator-Interface', '[repository withheld for review]')
    else:
        E.set_par(E.find_par(P, 'Weiskittel, A.R., Sprecher, I., Gottesman, A., Rice, B.', startswith=True), 'Authors withheld for double-blind review.')
    # deposit citation and DOI
    for p in P:
        t = p.text
        if t.startswith('[dataset] Weiskittel'):
            E.set_par(p, '[dataset] Authors, 2026. Data and prediction functions of this study, withheld for double-blind review. Zenodo. [DOI withheld for review]')
            continue
        if 'Weiskittel et al., 2026' in t:
            E.replace_in(p, 'Weiskittel et al., 2026', 'Authors, 2026', t.count('Weiskittel et al., 2026'))
        t = p.text
        if 'https://doi.org/10.5281/zenodo.22997704' in t:
            E.replace_in(p, 'https://doi.org/10.5281/zenodo.22997704', '[DOI withheld for review]', t.count('https://doi.org/10.5281/zenodo.22997704'))
        t = p.text
        if 'https://doi.org/10.5281/zenodo.21081014' in t:
            E.replace_in(p, 'https://doi.org/10.5281/zenodo.21081014', '[DOI withheld for review]', t.count('https://doi.org/10.5281/zenodo.21081014'))
    cp = d.core_properties
    cp.author = ''; cp.last_modified_by = ''; cp.title = ''; cp.comments = ''
    d.save(out)
    # leak check
    d2 = docx.Document(out)
    txt = '\n'.join(p.text for p in E.iter_all_pars(d2))
    leaks = [w for w in ['Weiskittel, A.R., Sprecher', 'Sprecher, I.', 'Gottesman', 'Ben Rice', 'Rice, B.', 'maine.edu', 'Forest Solutions', 'Midgard', 'zenodo.21081014', 'zenodo.22997704', 'Weiskittel et al., 2026'] if w in txt]
    print(out, 'leaks:', leaks)

run('/tmp/k6/out/koa_manuscript_v106_DRAFT.docx', '/tmp/k6/out/koa_manuscript_v106_DRAFT_BLIND.docx', 'ms')
run('/tmp/k6/out/koa_supplemental_v104.docx', '/tmp/k6/out/koa_supplemental_v104_BLIND.docx', 'sup')
