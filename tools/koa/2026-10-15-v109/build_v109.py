# -*- coding: utf-8 -*-
"""koa v109 build: manuscript (named and BLIND), title page and cover letter from the v108 package.

Applies red team 7 and stress test 4 fixes as scoped on 2026-10-15. Anchored edits only (docedit),
no equation is touched. Rerunnable: every output is rebuilt from the read-only inputs in /tmp/rt7/in.
"""
import sys, os, re, io, zipfile, base64, hashlib, subprocess, json, copy
sys.path.insert(0, '/tmp/v109/work')
import docx
import docedit as E
from counts import xml_counts

IN = '/tmp/rt7/in/'
OUT = '/tmp/v109/out/'
TMP = '/tmp/v109/work/tmp/'
os.makedirs(OUT, exist_ok=True); os.makedirs(TMP, exist_ok=True)
LOG = []

def log(*a):
    s = ' '.join(str(x) for x in a); LOG.append(s); print(s)

def rep(P, anchor, old, new, count=1):
    p = E.find_par(P, anchor)
    E.replace_in(p, old, new, count)
    return p

# ---------------------------------------------------------------- shared texts
FUND = ('This work was supported by the USDA Forest Service through a subcontract with Forest Solutions Inc., '
        'which commissioned the development of the equations. Two authors are employed by Forest Solutions Inc. and '
        'contributed to data curation, BYI modeling and spatial analysis as listed under Author contributions, and '
        'Forest Solutions Inc. set the rules for retaining records at duplicated and status-conflict plot visits '
        '(Section S1.3). The USDA Forest Service had no role in study design, in the collection, analysis or '
        'interpretation of data, in the writing of the report, or in the decision to submit the article for publication. '
        'Forest Solutions Inc. employees took part as authors in data curation and analysis, and the decision to submit '
        'was made by all authors.')

DATA1_OLD = 'the scripts that produced every table and figure (v103_rebuild_20260930.zip, engine_v103.zip) are archived'
DATA1_NEW = ('the scripts that produced every table and figure of the recomputed system (v103_rebuild_20260930.zip, '
             'engine_v103.zip), with the tables on the original covariates drawn from earlier versions of the deposit, '
             'are archived')
DATA1_TAIL_OLD = 'is the version of record for the results reported here.'
DATA1_TAIL_NEW = ('is the version of record for the results reported here. The deposit holds no plot coordinates, and '
                  'true Forest Inventory and Analysis plot coordinates are not released.')
DATA2_OLD = ('Because the variant and the simulator have differed during development in crown recession, in the stand '
             'basal area of the height increment density term and in the survival structure, projections from FVS-HI')
DATA2_NEW = ('Because the variant differs from the deposited simulator in crown recession, in the stand basal area of '
             'the height increment density term and in the survival structure, by up to 5.0 cm in DBH after ten years '
             'on a planted test list (Section S8), projections from FVS-HI')

# ---------------------------------------------------------------- manuscript
def edit_manuscript(src, out, blind):
    d = docx.Document(src)
    P = d.paragraphs
    A = list(E.iter_all_pars(d))
    if not blind:
        # S10 b: one form of each name
        rep(P, 'Irene Sprecherb, Aviva', 'Aaron Weiskittel', 'Aaron R. Weiskittel')
        rep(P, '*Corresponding author:', 'Aaron Weiskittel (', 'Aaron R. Weiskittel (')
        rep(P, 'Midgard Natural Resources, New Hampshire', 'Midgard Natural Resources,', 'Midgard Natural Resources LLC,')
        rep(P, 'Aaron Weiskittel: Conceptualization', 'Aaron Weiskittel:', 'Aaron R. Weiskittel:')
        # M3: funding
        E.set_par(E.find_par(P, 'This work was supported by the USDA', startswith=True), FUND)
    # Abstract: N11, S1, S2 (with the cuts that keep it at or below 300 words)
    ab = 'yet has no integrated growth and yield (GY) model'
    rep(P, ab, ' over seven decades.', '.')
    rep(P, ab, 'Biomass Yield Index (BYI), a productivity index derived from', 'Biomass Yield Index (BYI), derived from')
    rep(P, ab, 'mapped at 30 m by a random forest trained', 'mapped at 30 m by random forests trained')
    rep(P, ab, 'origin constants that also absorb differences among data sources,',
        'origin constants that also absorb data-source differences,')
    rep(P, ab, 'although the height increment site term was not distinguishable from zero.',
        'although the height increment site interval spans zero.')
    rep(P, ab, 'any mortality rises with stand density,',
        'any mortality rises with stand density in the short planted intervals that dominate the record,')
    rep(P, ab, 'from joint parameter draws, with installations resampled within data sources.', 'from joint parameter draws.')
    rep(P, ab, 'for management, economic and policy assessments, although projections for sites off Hawaii Island,',
        'for management and economic assessments, although projections beyond the 52-year record, for sites off Hawaii Island,')
    rep(P, ab, 'older than about 18 years rely on extrapolation.',
        'older than about 18 years rely on extrapolation, and planted volume is more likely overstated.')
    # Introduction: S15 (end of paragraph 4) and S6 (FVS-HI)
    rep(P, 'Comparable systems exist for data-limited tropical hardwoods', 'none has been developed for koa.',
        'none has been developed for koa. Beyond koa, the study tests whether three devices can carry an individual-tree '
        'model where stand age and long remeasurements are missing, namely an age-free productivity index, a mortality '
        'structure that separates occurrence from magnitude, and linked-system evaluation on remeasured plots.')
    rep(P, 'The goal of this analysis was to synthesize',
        'The equations were developed for the Hawaii variant of the Forest Vegetation Simulator (FVS-HI; Crookston and '
        'Dixon, 2005), which implements the height, crown and increment equations and can support further development for '
        'koa and extension to additional species.',
        'The equations were developed for a Hawaii variant of the Forest Vegetation Simulator (FVS-HI) built on the FVS '
        'framework (Crookston and Dixon, 2005). The variant implements the height, crown and increment equations but '
        'differs from the deposited simulator in crown recession, the height increment density term and the survival '
        'structure (Data and code availability).')
    # 2.2: S8 (first citation of Table 2 before Tables 3 and 4)
    rep(P, 'Table S1 defines every variable and its units.', 'Table S1 defines every variable and its units.',
        'Table S1 defines every variable and its units, and Table 2 summarizes them.')
    # 2.5: S16
    rep(P, 'The height and crown equations were cross-validated', 'Genuer et al., 2015) packages.',
        'Genuer et al., 2015) packages. Analysis and simulator code was written with assistance from Claude (Anthropic), '
        'and every reported table and figure was regenerated from the archived scripts (Section S8).')
    # 3.3.3: M6 and S9 b (ST4 D8)
    rep(P, 'Magnitude given occurrence showed no detectable density response',
        'the 5.6% yr−1 observed over the 282 Stage 1 intervals, 267 of them planted,',
        'the 5.6% {yr} observed over the 282 Stage 1 intervals, most of them planted,')
    # 3.5: S5
    rep(P, 'Uneven-aged natural stands, which carry ingrowth, reach', 'by age 100.',
        'by age 100. Planted stands peak at 57.2 {m2ha} and an SDI of 889 (Section S7.1), above the largest DOFAW basal '
        'area and the highest published koa basal area of 42 {m2ha} (Harrington et al., 1995; Section S7.2).')
    # 4.1: M4 a, b; S1; S4
    p = E.find_par(P, 'This study delivers the first linked individual-tree GY equations')
    E.replace_in(p, 'with installation-level support that is moderate for diameter increment and weak for height growth and crown base.',
                 'with installation-level support that is weak for every component, since only the diameter term fitted on '
                 'the original covariates has an interval that excludes zero.')
    E.replace_in(p, 'The mapped surface separates forest from nonforest and indicates only coarse productivity differences, '
                    'with demonstrated skill on Hawaii Island alone, and it can likely',
                 'The mapped surface separates forest from nonforest but has no demonstrated skill among forested plots, '
                 'even on Hawaii Island (Section 3.2), and it can likely')
    p = E.find_par(P, 'Koa mortality occurrence rose with stand density', startswith=True)
    E.replace_in(p, 'Koa mortality occurrence rose with stand density', 'In the planted record, mortality occurrence rose with stand density')
    E.replace_in(p, 'band in seven of nine stand types, and the two exceptions are the low natural sites, at −1.02 in even-aged '
                    'and −1.18 in uneven-aged stands.',
                 'band in seven of nine stand types at the default initial densities, and the two exceptions are the low '
                 'natural sites, at −1.02 in even-aged and −1.18 in uneven-aged stands, whereas from the sparsest starts at '
                 'BYI 264 the slopes are shallower (−0.81 natural, −0.93 planted).')
    # 4.4: S1, S4, S13
    p = E.find_par(P, 'Mortality was the component furthest from observation', startswith=True)
    E.replace_in(p, 'Ecologically, the pattern suggests that koa mortality arrives in episodes',
                 'If the planted pattern extends to natural stands, koa mortality would arrive in episodes')
    E.replace_in(p, 'The shallow realized slope on the low natural site (−1.02 over years 11 to 100), against the steeper slopes '
                    'of long-term European plots (Pretzsch and Biber, 2005; Pretzsch, 2006), most likely reflects the mortality level,',
                 'The shallow realized slope on the low natural site (−1.02 over years 11 to 100) lies within the −0.9 to −1.2 '
                 'range fitted to the upper boundary of the natural plot-measures (Section S7.2), although that boundary is '
                 'sparsely populated, and is shallower than the slopes of long-term European plots (Pretzsch and Biber, 2005; '
                 'Pretzsch, 2006), which may partly reflect the mortality level,')
    E.replace_in(p, 'observed stand density index (Section 3.3.3).', 'observed stand density index (Section 3.5).')
    # 4.5: M4 c
    p = E.find_par(P, 'For management, net MAI culminates early', startswith=True)
    E.replace_in(p, 'On low-productivity sites, natural stand growth is small', 'Where low productivity is established locally, natural stand growth is small')
    E.replace_in(p, 'On sites with BYI above 250,', 'At BYI levels above 250,')
    E.replace_in(p, 'which supports site screening but cannot underwrite',
                 'which supports comparison among BYI levels but not the screening of individual sites, since the mapped '
                 'surface cannot separate the three levels at a single location (Section 3.2), and cannot underwrite')
    # 4.6: same Table 5 age-60 defect as the note (S12 c, carried to the text)
    rep(P, 'Every row of Table 5 at age 60 and beyond is extrapolation', 'Every row of Table 5 at age 60 and beyond is extrapolation',
        'Every row of Table 5 at age 100, and of Table S30 at age 60 and beyond, is extrapolation')
    # Conclusions: S3, M4 d
    rep(P, 'Performance is mixed', 'and the natural mortality level needed its own calibration, which lowered the survival '
        'overprojection on the natural calibration intervals from +0.194 to +0.041.',
        'and the natural mortality level needed its own calibration, solved on the same natural intervals. On the 17 '
        'benchmark plots, all of which informed fitting or calibration, QMD was equivalent to observation in both origins, '
        'whereas natural survival and basal area in either origin were not.')
    rep(P, 'Net MAI culminates at 37, 25 and 20 years on the low, medium and high natural levels. At age 40',
        'is 160 to 197 m3 ha−1, or about 38 to 47 Mg C ha−1 of stemwood carbon,',
        'is 160 to 197 {m3ha} (95% Monte Carlo ranges 57 to 231 and 70 to 285), or about 38 to 47 {MgCha} of stemwood '
        'carbon (about 14 to 68), before the ±12% form-factor uncertainty,')
    rep(P, 'are a provisional foundation, and their first refinements', 'and a productivity surface that holds off Hawaii Island.',
        'and a productivity surface with demonstrated skill among forested plots, on and off Hawaii Island.')
    # Data and code availability: M5, S14, S6 (named and BLIND alike; the BLIND text keeps its withheld DOIs and URL)
    p = E.find_par(P, 'The assembled tree-level dataset (AK_TREE.csv', startswith=True)
    E.replace_in(p, DATA1_OLD, DATA1_NEW)
    E.replace_in(p, DATA1_TAIL_OLD, DATA1_TAIL_NEW)
    rep(P, 'The HiGy.R in the deposit reproduces', DATA2_OLD, DATA2_NEW)
    # Fig. 6 caption: S9 d, S12 b
    p = E.find_par(P, 'Behavior of the koa system at Biomass Yield Index 264')
    E.replace_in(p, 'stems per hectare (logarithmic axis)', 'stems {ha} (logarithmic axis)')
    E.replace_in(p, 'Breaks in the planted curves near age 35', 'Breaks in the planted curves at about age 35 to 40')
    # Table 5: M4 e (header, bold kept), S12 c, S5 and S9 c (note)
    E.replace_in(E.find_par(A, 'Site class (BYI)'), 'Site class (BYI)', '**BYI level**')
    p = E.find_par(P, 'Notes. Intervals are the 2.5th and 97.5th percentiles', startswith=True)
    E.replace_in(p, 'Rows at age 60 and beyond,', 'Rows at age 100,')
    E.replace_in(p, 'No point estimate in Table S30 exceeds 56.5 m2 ha−1 of basal area, and only the uneven-aged',
                 'Point estimates in Table S30 reach 56.5 {m2ha} of basal area, and the planted trajectory peaks at 57.2 '
                 '{m2ha} between grid ages (Section S7.1). Only the uneven-aged')
    if not blind:
        cp = d.core_properties
    d.save(out)
    return out

# ---------------------------------------------------------------- package surgery
CUSTOMXML_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml'

def surgery(path, strip_customxml, header_fix):
    zin = zipfile.ZipFile(path)
    items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    zin.close()
    out = []
    removed = []
    for info, data in items:
        n = info.filename
        if strip_customxml and n.startswith('customXml/'):
            removed.append(n); continue
        if strip_customxml and n == 'word/_rels/document.xml.rels':
            s = data.decode('utf-8')
            s2 = re.sub(r'<Relationship [^>]*Type="%s"[^>]*/>' % re.escape(CUSTOMXML_REL), '', s)
            if s2 != s: removed.append('rel customXml in ' + n)
            data = s2.encode('utf-8')
        if strip_customxml and n == '[Content_Types].xml':
            s = data.decode('utf-8')
            s2 = re.sub(r'<Override PartName="/customXml/[^"]*"[^>]*/>', '', s)
            if s2 != s: removed.append('override customXml in ' + n)
            data = s2.encode('utf-8')
        if header_fix and re.match(r'word/header\d+\.xml', n):
            s = data.decode('utf-8')
            s2 = s.replace('Weiskittel et al. – Koa G&amp;Y model', 'Weiskittel et al., Koa G&amp;Y model')
            if s2 != s: removed.append('running head fixed in ' + n)
            data = s2.encode('utf-8')
        if n == 'docProps/custom.xml':
            s = data.decode('utf-8')
            if re.search(r'<property', s):
                log('  docProps/custom.xml holds properties:', s[:300])
        out.append((info, data))
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w') as z:
        for info, data in out:
            z.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    os.replace(tmp, path)
    log('  surgery', os.path.basename(path), removed)

SWEEP = ['Weiskittel', 'Sprecher', 'Gottesman', 'Rice', 'Maine', 'Forest Solutions', 'Midgard', 'CRSF', 'orcid',
         'zenodo.21081014', 'zenodo.22997704', 'holoros', 'MidgardNaturalResources']

def sweep(path):
    hits = []
    z = zipfile.ZipFile(path)
    for n in z.namelist():
        data = z.read(n)
        texts = []
        if n.endswith('.xml') or n.endswith('.rels'):
            s = data.decode('utf-8', 'ignore'); texts.append(('xml', s))
            for m in re.finditer(r'roundtripDataSignature="[^"]*">([A-Za-z0-9+/=\s]+)<', s):
                try:
                    texts.append(('base64', base64.b64decode(m.group(1)).decode('utf-8', 'ignore')))
                except Exception as e:
                    texts.append(('base64-error', str(e)))
        else:
            texts.append(('binary', data.decode('latin-1')))
        for kind, s in texts:
            plain = re.sub(r'<[^>]+>', '', s) if kind == 'xml' else s
            for pat in SWEEP:
                flags = re.I if pat == 'orcid' else 0
                for m in re.finditer(r'\b' + re.escape(pat) if pat in ('Rice', 'Maine') else re.escape(pat), plain, flags):
                    ctx = plain[max(0, m.start() - 60):m.end() + 60].replace('\n', ' ')
                    allowed = pat == 'Sprecher' and 'Owen, S.M.' in plain[max(0, m.start() - 200):m.start()]
                    near = plain[max(0, m.start() - 40):m.end() + 40]
                    if pat == 'Weiskittel' and re.search(r'Weiskittel( et al\., 2011|, A\.R\., (Hann|McGarrigle|Allen|Kershaw|2020)|, A\., 2019| \(2020\)|, 2020\))', near):
                        allowed = 'third-person literature citation'
                    hits.append(dict(part=n, kind=kind, pattern=pat, allowed=allowed, context=ctx))
    return hits

# ---------------------------------------------------------------- title page
def sup_at(p, sub, offset):
    model = E.chars(p)
    text = ''.join(c for c, _ in model)
    i = text.index(sub) + offset
    c, r = model[i]
    nr = copy.deepcopy(r) if r is not None else E.etree.Element(E.qn('w:rPr'))
    E._set_flag(nr, 'w:vertAlign', True, 'superscript')
    model[i] = [c, nr]
    E.write_chars(p, model)

def build_title_page(named_ms):
    d = docx.Document(IN + 'koa_title_page_v108.docx'); P = d.paragraphs
    p = E.find_par(P, 'Irene Sprecherb, Aviva')
    assert 'Weiskittela,*,' in p.text
    sup_at(p, 'Weiskittela,*,', len('Weiskittela,'))      # S10 c: a,* fully superscript
    rep(P, 'Midgard Natural Resources, New Hampshire', 'Midgard Natural Resources,', 'Midgard Natural Resources LLC,')
    E.set_par(E.find_par(P, 'This work was supported by the USDA', startswith=True), FUND)
    # S10 d: the data statement is the manuscript text, both paragraphs
    m = docx.Document(named_ms); MP = m.paragraphs
    t1 = E.find_par(MP, 'The assembled tree-level dataset (AK_TREE.csv', startswith=True).text
    t2 = E.find_par(MP, 'The HiGy.R in the deposit reproduces', startswith=True).text
    p = E.find_par(P, 'The assembled tree-level dataset', startswith=True)
    E.set_par(p, t1)
    p2 = copy.deepcopy(p._p); p._p.addnext(p2)
    from docx.text.paragraph import Paragraph
    q = Paragraph(p2, p._parent)
    E.set_par(q, t2.replace('1.4 × 10−15', '1.4 × 10^−15^'))
    for q in d.paragraphs:   # keep each section heading with its text (render check)
        if q.text.strip() in ('Article type', 'Author contributions (CRediT)', 'Acknowledgments', 'Funding',
                              'Declaration of competing interest', 'Ethics approval and consent to participate',
                              'Data availability',
                              'Declaration of generative AI and AI-assisted technologies in the writing process'):
            q.paragraph_format.keep_with_next = True
    out = OUT + 'koa_title_page_v109.docx'
    d.save(out)
    return out

# ---------------------------------------------------------------- cover letter
NUMW = {5: 'five', 6: 'six', 8: 'eight', 33: '33', 62: '62'}

def build_cover(c):
    d = docx.Document(IN + 'koa_cover_letter_v107_FINAL.docx'); P = d.paragraphs
    E.set_par(E.find_par(P, 'October 1, 2026', startswith=True), 'October 15, 2026')
    rep(P, 'Please consider the enclosed manuscript', '*A. koa* is the most valuable'.replace('*', ''),
        'Koa (*A. koa*) is the most valuable')
    p = E.find_par(P, 'Three contributions should interest', startswith=True)
    E.replace_in(p, 'mapped at 30 m by random forests trained on Hawaii Island.',
                 'mapped at 30 m by random forests trained on Hawaii Island, which separates forest from nonforest but not yet '
                 'productivity among forested plots.')
    E.replace_in(p, 'records any mortality rises with stand density, whereas',
                 'records any mortality rises with stand density in the short planted intervals that dominate the record, whereas')
    E.replace_in(p, 'and density and thinning scenarios.',
                 'and density and thinning scenarios. Beyond koa, the study tests whether these three devices, an age-free '
                 'productivity index, a mortality structure that separates occurrence from magnitude and linked-system '
                 'evaluation on remeasured plots, can carry an individual-tree model where stand age and long remeasurements '
                 'are missing.')
    E.replace_in(p, 'The Discussion states the resulting rotation and site-class guidance together with its limits, since',
                 'The Discussion states what the projections can and cannot support for rotation and carbon decisions, since')
    p = E.find_par(P, 'Every projection in the paper was generated', startswith=True)
    old_first = p.text.split(' The main text runs')[0]
    E.replace_in(p, old_first,
                 'Every projection in the paper was generated under the height equation of Table 3, a height increment '
                 'equation refitted on the recomputed stand covariates, a diameter increment vector fitted on the original '
                 'covariates whose origin constants were re-solved on the recomputed frame, the natural mortality level '
                 'calibrated on 24 remeasured natural plot intervals and a survival equation respecified on a recovered sample '
                 'carrying 877 deaths, with every interval drawn jointly from the parameters and from installations resampled '
                 'within data sources.')
    words = int(round(c['body'] / 100.0) * 100)
    E.replace_in(p, 'The main text runs about 8,400 words with a 299-word abstract, eight keywords, five tables, six figures and '
                    '62 references, with a supplement of 33 tables and eight figures.',
                 'The main text runs about {:,} words with a {}-word abstract, {} keywords, {} tables, {} figures and {} '
                 'references, with a supplement of {} tables and {} figures.'.format(
                     words, c['abstract'], NUMW[c['keywords']], NUMW[c['tables']], NUMW[c['figs']], c['refs'],
                     NUMW[c['s_tables']], NUMW[c['s_figs']]))
    E.replace_in(p, 'with fuzzed plot positions,', 'with no plot coordinates,')
    E.replace_in(E.find_par(P, 'The manuscript is original', startswith=True),
                 'All authors have approved the submission and declare no competing interests.',
                 'All authors have approved the submission. Competing interests are declared in the manuscript and on the '
                 'title page: two authors are employees of Forest Solutions Inc., which commissioned this work and uses the '
                 'resulting yield curves in a koa financial model, and one author is affiliated with Midgard Natural Resources '
                 'LLC, which holds the scope of work under which the equations were developed.')
    # keep the letter on one page: the v109 text is about six lines longer than v107
    from docx.shared import Twips, Pt
    s = d.sections[0]
    s.top_margin = s.bottom_margin = Twips(720); s.left_margin = s.right_margin = Twips(1224)
    d.styles['Normal'].font.size = Pt(11.5)
    out = OUT + 'koa_cover_letter_v109.docx'
    d.save(out)
    return out

# ---------------------------------------------------------------- counts
def ms_counts(path):
    d = docx.Document(path); T = [p.text.strip() for p in d.paragraphs]
    abstract = T[T.index('Abstract') + 1]
    refs = [t for t in T[T.index('References') + 1:T.index('Figures')] if t]
    kw = [t for t in T if t.startswith('Keywords:')][0]
    return dict(abstract=len(abstract.split()), body=xml_counts(path)['xml_intro_to_concl'], refs=len(refs),
                keywords=len(kw.split(':', 1)[1].split(';')), tables=len(d.tables),
                figs=sum(1 for t in T if re.match(r'Fig\. \d+\.', t)),
                table_captions=sum(1 for t in T if re.match(r'Table \d+\.', t)))

def supp_counts():
    s = open(IN + 'koa_supplemental_v104.md', encoding='utf-8').read()
    return dict(s_tables=len(set(re.findall(r'^\*\*Table S(\d+)\.\*\*', s, re.M))),
                s_figs=len(set(re.findall(r'^\*\*(?:Fig\.|Figure) S(\d+)\.\*\*', s, re.M))))

def letter_counts(path):
    t = '\n'.join(p.text for p in docx.Document(path).paragraphs)
    m = re.search(r'about ([\d,]+) words with a (\d+)-word abstract, (\w+) keywords, (\w+) tables, (\w+) figures and (\d+) '
                  r'references, with a supplement of (\d+) tables and (\w+) figures', t)
    return m.groups() if m else None

def md5(p):
    return hashlib.md5(open(p, 'rb').read()).hexdigest()

def text_of(path):
    d = docx.Document(path)
    return '\n'.join(p.text for p in E.iter_all_pars(d))

# ---------------------------------------------------------------- render
def render(docx_path):
    env = dict(os.environ, HOME='/tmp/lohome', TMPDIR='/tmp', XDG_CACHE_HOME='/tmp/lohome/c', XDG_CONFIG_HOME='/tmp/lohome/cfg')
    os.makedirs('/tmp/lohome/c', exist_ok=True); os.makedirs('/tmp/lohome/cfg', exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', TMP, docx_path], env=env, check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    src = TMP + os.path.basename(docx_path)[:-5] + '.pdf'
    dst = OUT + os.path.basename(docx_path)[:-5] + '_REVIEW.pdf'
    os.replace(src, dst)
    pages = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', dst], capture_output=True, text=True).stdout).group(1))
    return dst, pages

# ---------------------------------------------------------------- main
if __name__ == '__main__':
    named = edit_manuscript(IN + 'koa_manuscript_v108_DRAFT.docx', OUT + 'koa_manuscript_v109_DRAFT.docx', False)
    blind = edit_manuscript(IN + 'koa_manuscript_v108_DRAFT_BLIND.docx', OUT + 'koa_manuscript_v109_DRAFT_BLIND.docx', True)
    surgery(named, strip_customxml=True, header_fix=True)
    surgery(blind, strip_customxml=True, header_fix=False)
    tp = build_title_page(named)
    surgery(tp, strip_customxml=False, header_fix=False)
    c_old = ms_counts(IN + 'koa_manuscript_v108_DRAFT.docx'); c_new = ms_counts(named); c_blind = ms_counts(blind)
    c_new.update(supp_counts())
    log('counts v108', c_old); log('counts v109', c_new); log('counts v109 BLIND', c_blind)
    assert c_new['abstract'] <= 300, c_new['abstract']
    assert c_blind['abstract'] == c_new['abstract']
    cl = build_cover(c_new)
    log('cover letter counts', letter_counts(cl))
    # checks
    res = {'counts_v108': c_old, 'counts_v109': c_new, 'counts_v109_blind': c_blind, 'letter': letter_counts(cl)}
    for f in [named, blind, tp, cl]:
        t = text_of(f)
        log(os.path.basename(f), 'em dashes:', t.count('—'), 'spaced en dash:', len(re.findall(' – ', t)))
    lt = text_of(cl); mt = text_of(named)
    for s in ['9,164', '158 installations', '17 plots', '15 plots', '52 years', '3.9 cm', '877 deaths', '24 remeasured natural plot intervals']:
        key = {'17 plots': '17 remeasured plots', '15 plots': '15 plots', '52 years': '52 years', '158 installations': '158 installations',
               '24 remeasured natural plot intervals': '24 natural plot intervals', '877 deaths': '877 deaths'}.get(s, s)
        log('  letter number', s, 'in letter:', s in lt, '| manuscript has', repr(key), key in mt)
    res['sweep'] = {}
    for f in [blind, named]:
        hits = sweep(f)
        res['sweep'][os.path.basename(f)] = hits
        log('sweep', os.path.basename(f), len(hits), 'hits;', sum(1 for h in hits if not h['allowed']), 'not allowed')
    pdfs = {}
    for f in [named, blind, tp, cl]:
        pdf, n = render(f); pdfs[os.path.basename(f)] = (os.path.basename(pdf), n); log('rendered', pdf, n, 'pages')
    res['pdfs'] = pdfs
    res['md5'] = {os.path.basename(p): md5(OUT + p) for p in sorted(os.listdir(OUT)) if p.endswith(('.docx', '.pdf'))}
    for k, v in res['md5'].items(): log('md5', k, v)
    json.dump(res, open(TMP + 'build_v109_result.json', 'w'), indent=1, ensure_ascii=False)

# Post-build edit applied 2026-10-15 in the main session: first citation of Fig. 4 in 2.6 (guide numbering rule)
import sys as _s; _s.path.insert(0, '/tmp/v109/work')
import docx as _d, docedit as _E
for _f in ['koa_manuscript_v109_DRAFT', 'koa_manuscript_v109_DRAFT_BLIND']:
    _p = '/tmp/v109/out/%s.docx' % _f
    _doc = _d.Document(_p)
    _q = _E.find_par(_doc.paragraphs, 'come from plot-cluster bootstraps')
    _E.replace_in(_q, 'come from plot-cluster bootstraps.', 'come from plot-cluster bootstraps, and the resulting bands are drawn with the projections (Fig. 4).')
    _doc.save(_p)

# Post-build edit 2026-10-15 (Aaron): international log pricing in $ per cubic meter
for _f in ['koa_manuscript_v109_DRAFT', 'koa_manuscript_v109_DRAFT_BLIND']:
    _p = '/tmp/v109/out/%s.docx' % _f
    _doc = _d.Document(_p)
    _q = _E.find_par(_doc.paragraphs, 'exceeding US$50 per board foot')
    _E.replace_in(_q, 'exceeding US$50 per board foot', 'exceeding about US$21,000 m^−3^ (US$50 per board foot)')
    _doc.save(_p)
