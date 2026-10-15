# -*- coding: utf-8 -*-
"""v108 from v107 FINAL: funding, competing interest, FVS-HI paragraph. Anchored edits only."""
import sys; sys.path.insert(0, '/tmp/k8/work')
import docx, docedit as E
SRC = '/mnt/user-data/uploads/2026--Koa/submission_v102/'
FUND = ('This work was supported by the USDA Forest Service through a subcontract with Forest Solutions Inc., '
        'which commissioned the development of the equations. Two authors are employed by Forest Solutions Inc. and '
        'contributed to data curation, BYI modeling and spatial analysis as listed under Author contributions, and '
        'Forest Solutions Inc. set the rules for retaining records at duplicated and status-conflict plot visits. '
        'The USDA Forest Service had no role in study design, in the collection, analysis or interpretation of data, '
        'in the writing of the report, or in the decision to submit the article for publication.')
COI = ('I. Sprecher and A. Gottesman are employees of Forest Solutions Inc., which commissioned this work and uses the '
       'resulting yield curves in a koa financial model. B. Rice is affiliated with Midgard Natural Resources LLC, which '
       'holds the scope of work with Forest Solutions Inc. under which the equations were developed and implemented in '
       'FVS-HI. A.R. Weiskittel declares no competing interests.')
def fvs(repo):
    return ('The HiGy.R in the deposit reproduces the expansion-weighted live-list BAL of Section 2.6 to 1.4 × 10^−15^ '
            'on the BAL test cases, and the FVS-HI variant built on it is maintained at ' + repo + '. Because the variant '
            'and the simulator have differed during development in crown recession, in the stand basal area of the height '
            'increment density term and in the survival structure, projections from FVS-HI should be checked against the '
            'deposited simulator (koa_projector.py), including the three-stage mortality of Section 2.6, before use.')
def edit(src, out, blind):
    d = docx.Document(src); P = d.paragraphs
    if not blind:
        E.set_par(E.find_par(P, 'This work was supported by the USDA', startswith=True), FUND)
        E.set_par(E.find_par(P, 'The authors declare no competing interests.', startswith=True), COI)
        E.set_par(E.find_par(P, 'The FVS-HI variant (HiGy.R,', startswith=True),
                  fvs('https://github.com/MidgardNaturalResources/ForestVegetationSimulator-Interface'))
    else:
        E.set_par(E.find_par(P, 'The authors declare no competing interests.', startswith=True), 'Withheld for double-blind review.')
        E.set_par(E.find_par(P, 'The FVS-HI variant (HiGy.R,', startswith=True), fvs('[repository withheld for review]'))
    d.save(out)
    txt = '\n'.join(p.text for p in E.iter_all_pars(docx.Document(out)))
    if blind:
        leaks = [w for w in ['Sprecher', 'Gottesman', 'Ben Rice', 'Rice, B.', 'B. Rice', 'maine.edu', 'Forest Solutions', 'Midgard', 'zenodo.21081014', 'zenodo.22997704', 'Weiskittel et al., 2026'] if w in txt]
        print(out, 'leaks:', leaks)
    print(out, 'ok, —:', txt.count('—'))
edit(SRC + 'koa_manuscript_v107_FINAL.docx', '/tmp/k8/out/koa_manuscript_v108_DRAFT.docx', False)
edit(SRC + 'koa_manuscript_v107_FINAL_BLIND.docx', '/tmp/k8/out/koa_manuscript_v108_DRAFT_BLIND.docx', True)
# title page
d = docx.Document(SRC + 'koa_title_page_v106.docx'); P = d.paragraphs
E.set_par(E.find_par(P, 'This work was supported by the USDA', startswith=True), FUND)
E.set_par(E.find_par(P, 'The authors declare no competing interests.', startswith=True), COI)
d.save('/tmp/k8/out/koa_title_page_v108.docx'); print('title page ok')
