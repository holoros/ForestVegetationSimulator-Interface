import sys, re; sys.path.insert(0,'/tmp/v109/work')
import docx, docedit as E
def counts(path):
    d=docx.Document(path); P=d.paragraphs
    T=[p.text.strip() for p in P]
    i_abs=T.index('Abstract'); abstract=T[i_abs+1]
    i0=T.index('1 Introduction'); i5=T.index('5 Conclusions'); iack=T.index('Acknowledgments'); 
    i_disc_end=i5
    body=[t for t in T[i0:iack] if t]
    body_nohead=[t for p,t in zip(P[i0:iack],T[i0:iack]) if t and not p.style.name.startswith('Heading')]
    iref=T.index('References'); ifig=T.index('Figures')
    refs=[t for t in T[iref+1:ifig] if t]
    nw=lambda L: sum(len(x.split()) for x in L)
    kw=[t for t in T if t.startswith('Keywords:')][0]
    return dict(abstract=len(abstract.split()), body_with_headings=nw(body), body_no_headings=nw(body_nohead),
                intro_to_disc=nw([t for t in T[i0:i5] if t]), refs=len(refs), keywords=len(kw.split(':',1)[1].split(';')),
                tables=len(d.tables), figs=sum(1 for t in T if re.match(r'Fig\. \d+\.',t)),
                tab_caps=sum(1 for t in T if re.match(r'Table \d+\.',t)))
if __name__=='__main__':
    for f in sys.argv[1:]: print(f, counts(f))

def xml_counts(path):
    d=docx.Document(path)
    W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    pars=[]
    for p in d.element.body.iter(W+'p'):
        # skip paragraphs inside tables
        anc=p.getparent(); intable=False
        while anc is not None:
            if anc.tag==W+'tbl': intable=True;break
            anc=anc.getparent()
        if intable: continue
        t=''.join(x.text or '' for x in p.iter(W+'t'))
        pars.append(t.strip())
    i0=pars.index('1 Introduction'); iack=pars.index('Acknowledgments'); i5=pars.index('5 Conclusions')
    nw=lambda L: sum(len(x.split()) for x in L)
    return dict(xml_intro_to_concl=nw(pars[i0:iack]), xml_intro_to_disc=nw(pars[i0:i5]))
if __name__=='__main__':
    for f in sys.argv[1:]: print(f, xml_counts(f))
